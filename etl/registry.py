import openpyxl, re, json, collections
import os; P=os.path.join(os.path.dirname(os.path.abspath(__file__)),'..','data','Dock_Schedule_-_Synthetic_Sample.xlsx')
wb=openpyxl.load_workbook(P,data_only=True)
PREF=r"(?:R/V|M/V|M/Y|S/V|S/Y|F/V|OS/V|OSV|Tug|Barge)"
ORGS=set()
def canon(s):
    s=re.sub(r'\s+',' ',s.strip()); s=re.sub(r'^OS/V\b','OSV',s,flags=re.I)
    m=re.match(r'^('+PREF+r")\s+(.*)$",s,re.I)
    if not m: return None
    pre=m.group(1).upper(); pre={'OSV':'OSV','TUG':'Tug','BARGE':'Barge'}.get(pre,pre)
    rest=re.sub(r"\s*\d{2,3}\s*'\s*$",'',m.group(2)).strip()
    return pre+' '+' '.join(w.capitalize() for w in rest.split())

vessels={}
for sh in ['Science','Yachts']:
    ws=wb[sh]
    cur=None
    for r in range(1,ws.max_row+1):
        rowvals=[(c,ws.cell(r,c).value) for c in range(1,ws.max_column+1)]
        # find a vessel name anywhere in this row
        anchor=None
        for c,v in rowvals:
            if isinstance(v,str) and re.match(PREF+r"\b",v.strip(),re.I):
                nm=canon(v)
                if nm:
                    anchor=(nm,c)
                    m=re.search(r"(\d{2,3})\s*'",v)
                    rec=vessels.setdefault(nm,dict(name=nm,loa=None,draft=None,operator=None,
                        contacts=[],emails=[],notes=[],url=None,source=sh))
                    if m and not rec['loa']: rec['loa']=int(m.group(1))
                    break
        if anchor: cur=anchor[0]
        if not cur: continue
        rec=vessels[cur]
        for c,v in rowvals:
            if not isinstance(v,str): continue
            v=v.strip()
            if not v or re.match(PREF+r"\b",v,re.I): continue
            m=re.search(r"LOA:\s*(\d{2,3})\s*'",v)
            if m:
                rec['loa']=rec['loa'] or int(m.group(1))
                dm=re.search(r"Draft:\s*(\d{1,2})\s*'",v)
                if dm: rec['draft']=int(dm.group(1))
                continue
            if '@' in v and '.' in v: rec['emails'].append(v); continue
            if re.match(r'^(Cell|Work|Ph|Tel)\s*[:#]',v,re.I): rec['contacts'].append(v); continue
            if re.match(r'^(Capt\.|Capt )',v) or (len(v.split())==2 and v[0].isupper() and v.split()[1][0].isupper() and not any(ch.isdigit() for ch in v)):
                rec['contacts'].append(v); continue
            if v.startswith('http'): rec['url']=v; continue
            if v.upper() in ('VESSEL','OPERATOR','CONTACT','WORK#','CELL#','EMAIL','NOTES'): continue
            if re.search(r'(University|Institute|Agency|Partners|Trust|Charters|Group|School|Foundation|Offshore|Services|Academy)',v):
                rec['operator']=rec['operator'] or v; ORGS.add(v); continue
            rec['notes'].append(v)
print('registry vessels:',len(vessels))
print('with LOA:',sum(1 for v in vessels.values() if v['loa']))
print('with draft:',sum(1 for v in vessels.values() if v['draft']))
print('with operator:',sum(1 for v in vessels.values() if v['operator']))
print('orgs:',sorted(ORGS))
import statistics
by=collections.defaultdict(list)
for v in vessels.values():
    if v['loa']: by[v['name'].split()[0]].append(v['loa'])
print()
for k,vv in sorted(by.items()): print(f'  {k:6s} n={len(vv):3d} min={min(vv):3d} med={int(statistics.median(vv)):3d} max={max(vv):3d}')
json.dump(list(vessels.values()),open('registry.json','w'),indent=0)
